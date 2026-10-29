"""प्रतिवेदन मेनु: उपभोक्ता समिति, बोलपत्र ठेक्का र दरभाउपत्र खरिद — एकीकृत वा विधिअनुसार छुट्टै;
योजनागत वित्तीय विवरण, कर कट्टी, भौतिक उपलब्धि — प्रिन्ट र Excel सहित."""
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect, render
from django.urls import reverse

from accounts.models import Ward
from accounts.permissions import has_perm
from accounts.wards import sorted_wards
from projects.models import FiscalYear, Sector, WorkCategory

from projects.models import Project

from .reports import KIND_PTYPE, KINDS, STATUS, all_rows, by_kind, by_sector, by_tax_head, by_ward, totals
from .views_workflow import _user_wards


def _filtered(request):
    """(rows, ctx) — फिल्टर: खरिद विधि, आ.व. (default: पछिल्लो), वडा, स्थिति, कामको प्रकार."""
    qs = Project.objects.filter(procurement_type_id__in=KIND_PTYPE.values())
    my_wards = _user_wards(request.user)
    if my_wards is not None:
        qs = qs.filter(ward_id__in=my_wards)

    fys = list(FiscalYear.objects.all().order_by("-id"))
    fy = request.GET.get("fiscalyear")
    if fy is None:            # पहिलो पटक: पछिल्लो आर्थिक वर्ष
        fy = str(fys[0].id) if fys else "all"
    ward = request.GET.get("ward", "")
    status = request.GET.get("status", "")
    cat = request.GET.get("category", "")
    kind = request.GET.get("kind", "")
    kind = kind if kind in KINDS else ""
    sector = request.GET.get("sector", "")
    if fy.isdigit():
        qs = qs.filter(fiscalyear_id=fy)
    if ward.isdigit():
        qs = qs.filter(ward_id=ward)
    if cat.isdigit():
        qs = qs.filter(work_category_id=cat)
    if sector.isdigit():
        qs = qs.filter(sector_id=sector)
    elif sector == "none":
        qs = qs.filter(sector__isnull=True)
    status = status if status in STATUS else ""

    all_wards = sorted_wards(Ward.objects.filter(id__in=my_wards) if my_wards is not None else None)
    wards = [w for w in all_wards if not ward.isdigit() or str(w.id) == ward]
    rows = all_rows(list(qs.values_list("id", flat=True)), kind, status)

    fy_name = next((x.name for x in fys if str(x.id) == fy), "")
    desc = [KINDS[kind] if kind else "उपभोक्ता समिति, बोलपत्र र दरभाउपत्र"]
    desc.append(f"आ.व. {fy_name}" if fy_name else "सबै आर्थिक वर्ष")
    desc.append(f"वडा नं. {wards[0].name}" if ward.isdigit() and wards else "सबै वडा")
    if status:
        desc.append(STATUS[status])
    if sector.isdigit():
        sc = Sector.objects.filter(id=sector).first()
        if sc:
            desc.append(sc.name)
    if cat.isdigit():
        c = WorkCategory.objects.filter(id=cat).first()
        if c:
            desc.append(c.name)
    keep = request.GET.copy()
    keep.pop("export", None)
    keep["fiscalyear"] = fy
    ctx = {
        "fiscal_years": fys, "all_wards": all_wards, "wards": wards, "statuses": STATUS.items(),
        "categories": WorkCategory.objects.all(), "kinds": KINDS.items(), "kind_label": KINDS.get(kind, ""),
        "sectors": Sector.objects.all(),
        "f": {"fiscalyear": fy, "ward": ward, "status": status, "category": cat, "kind": kind, "sector": sector},
        "filter_desc": " · ".join(desc), "qs": keep.urlencode(),
        "report_tabs": [
            ("summary", "सारांश", reverse("uc_report_summary"), "report.uc"),
            ("projects", "योजनागत वित्तीय विवरण", reverse("uc_report_projects"), "report.uc"),
            ("taxes", "कर कट्टी", reverse("uc_report_taxes"), "report.uc"),
            ("outputs", "भौतिक उपलब्धि", reverse("outputs_report") + f"?fiscalyear={fy if fy.isdigit() else ''}&ward={ward}",
             "report.outputs"),
        ],
    }
    return rows, ctx


def _title(ctx, name):
    return f"{ctx['kind_label']} — {name}" if ctx["kind_label"] else f"खरिद तथा भुक्तानी — {name}"


def _guard(request):
    return has_perm(request.user, "report.uc")


@login_required
def summary(request):
    if not _guard(request):
        return redirect("unauthorized")
    rows, ctx = _filtered(request)
    if request.GET.get("export") == "xlsx":
        return _xlsx(rows, ctx)
    from projects.models import Project
    from projects.outputs import report as outputs_report

    ctx.update({
        "t": totals(rows),
        "kinds_table": by_kind(rows),
        "sectors_table": by_sector(rows),
        "wards_table": by_ward(rows, ctx["wards"]),
        "tax_heads": by_tax_head(rows, ctx["wards"]),
        "outputs": outputs_report(Project.objects.filter(id__in=[r["p"].id for r in rows]), ctx["wards"]),
        "status_labels": STATUS,
    })
    return render(request, "reports/uc_summary.html", ctx | {"report_title": _title(ctx, "एकीकृत प्रतिवेदन")})


@login_required
def projects(request):
    if not _guard(request):
        return redirect("unauthorized")
    rows, ctx = _filtered(request)
    if request.GET.get("export") == "xlsx":
        return _xlsx(rows, ctx)
    ctx.update({"rows": rows, "t": totals(rows)})
    return render(request, "reports/uc_projects.html", ctx | {"report_title": _title(ctx, "योजनागत वित्तीय विवरण")})


@login_required
def taxes(request):
    if not _guard(request):
        return redirect("unauthorized")
    rows, ctx = _filtered(request)
    if request.GET.get("export") == "xlsx":
        return _xlsx(rows, ctx)
    heads = by_tax_head(rows, ctx["wards"])
    ctx.update({"heads": heads, "t": totals(rows)})
    return render(request, "reports/uc_tax.html", ctx | {"report_title": _title(ctx, "कर कट्टी प्रतिवेदन")})


# ---------------------------------------------------------------- Excel (सबै पाना एउटै फाइलमा)
def _xlsx(rows, ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from projects.models import Project
    from projects.outputs import report as outputs_report

    wards = ctx["wards"]
    t = totals(rows)
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="E8EEFB")

    def sheet(ws, header, data, widths):
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head_fill
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in data:
            ws.append(row)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        ws.freeze_panes = "A2"

    f = float
    ws = wb.active
    ws.title = "सारांश"
    ws.append([_title(ctx, "एकीकृत प्रतिवेदन")])
    ws.append([ctx["filter_desc"]])
    ws.append([])
    for label, val in [("कुल योजना", t["count"])] + [(STATUS[k], v) for k, v in t["by_status"].items()] + [
            ("विनियोजित रकम", f(t["allocated"])), ("लागत अनुमान (कुल)", f(t["estimate"])),
            ("सम्झौता रकम", f(t["contract"])), ("उपभोक्ता अंशदान (लागत अनुमान)", f(t["user_share"])),
            ("मूल्याङ्कन रकम", f(t["valuation"])), ("कुल भुक्तानी", f(t["gross"])), ("कर कट्टी", f(t["tax"])),
            ("खुद भुक्तानी", f(t["net"])), ("बाँकी विनियोजन", f(t["remaining"]))]:
        ws.append([label, val])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 18

    sheet(wb.create_sheet("विधिअनुसार"),
          ["खरिद विधि", "योजना", "सम्पन्न", "विनियोजित", "लागत अनुमान", "सम्झौता", "कुल भुक्तानी", "कर कट्टी",
           "खुद भुक्तानी", "बाँकी"],
          [[k["label"], k["count"], k["by_status"]["completed"], f(k["allocated"]), f(k["estimate"]), f(k["contract"]),
            f(k["gross"]), f(k["tax"]), f(k["net"]), f(k["remaining"])] for k in by_kind(rows)],
          [18, 8, 9, 15, 15, 15, 15, 13, 15, 15])

    sheet(wb.create_sheet("क्षेत्रअनुसार"),
          ["क्षेत्र", "योजना", "सम्पन्न", "विनियोजित", "लागत अनुमान", "सम्झौता", "कुल भुक्तानी", "कर कट्टी", "खुद भुक्तानी", "बाँकी"],
          [[k["name"], k["count"], k["by_status"]["completed"], f(k["allocated"]), f(k["estimate"]), f(k["contract"]),
            f(k["gross"]), f(k["tax"]), f(k["net"]), f(k["remaining"])] for k in by_sector(rows)],
          [30, 8, 9, 15, 15, 15, 15, 13, 15, 15])

    sheet(wb.create_sheet("वडागत"),
          ["वडा", "योजना", "सम्पन्न", "विनियोजित", "लागत अनुमान", "कुल भुक्तानी", "कर कट्टी", "खुद भुक्तानी", "बाँकी"],
          [[(w["ward"].name if w["ward"] else "अन्य"), w["count"], w["by_status"]["completed"], f(w["allocated"]),
            f(w["estimate"]), f(w["gross"]), f(w["tax"]), f(w["net"]), f(w["remaining"])] for w in by_ward(rows, wards)],
          [8, 8, 9, 15, 15, 15, 13, 15, 15])

    sheet(wb.create_sheet("योजनागत"),
          ["योजना", "वडा", "क्षेत्र", "उपक्षेत्र", "खरिद विधि", "उपभोक्ता समिति / निर्माण व्यवसायी / फर्म", "कामको प्रकार", "स्थिति", "विनियोजित",
           "लागत अनुमान", "सम्झौता रकम", "उपभोक्ता अंशदान", "मूल्याङ्कन", "किस्ता / अग्रिम / रनिङ", "कुल भुक्तानी",
           "कर कट्टी", "खुद भुक्तानी", "बाँकी"],
          [[r["p"].name, getattr(r["p"].ward, "name", ""), getattr(r["p"].sector, "name", ""),
            getattr(r["p"].subsector, "name", ""), r["kind_label"], r["party"], r["work_category"], r["status_label"],
            f(r["allocated"]), f(r["estimate"]) if r["estimate"] is not None else None,
            f(r["contract"]) if r.get("contract") is not None else None, f(r["user_share"]),
            f(r["valuation"]) if r["valuation"] is not None else None, f(r["installments"]), f(r["gross"]),
            f(r["tax"]), f(r["net"]), f(r["remaining"])] for r in rows],
          [36, 6, 22, 22, 14, 28, 14, 16] + [14] * 10)

    heads = by_tax_head(rows, wards)
    sheet(wb.create_sheet("कर कट्टी"),
          ["कर शीर्षक", "जम्मा", "योजना"] + [f"वडा {w.name}" for w in wards],
          [[h["name"], f(h["total"]), h["count"]] + [f(v) if v else None for v in h["ward_values"]] for h in heads],
          [28, 15, 8] + [12] * len(wards))

    groups = outputs_report(Project.objects.filter(id__in=[r["p"].id for r in rows]), wards)
    sheet(wb.create_sheet("भौतिक उपलब्धि"),
          ["कामको प्रकार", "उपलब्धि", "एकाइ", "जम्मा", "योजना"] + [f"वडा {w.name}" for w in wards],
          [[g["cat"].name, r["ind"].name, r["ind"].unit, f(r["total"]), r["count"]]
           + [f(v) if v else None for v in r["ward_values"]] for g in groups for r in g["rows"]],
          [16, 32, 10, 12, 8] + [10] * len(wards))

    resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = 'attachment; filename="kharid-bhuktani-pratibedan.xlsx"'
    wb.save(resp)
    return resp
