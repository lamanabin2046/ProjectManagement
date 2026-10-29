"""बजेट नियन्त्रण: बजेट शीर्षक (सेटिङ), बजेट थप / कटौती / रकमान्तर, र बजेट स्थिति प्रतिवेदन."""
from collections import OrderedDict
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db import transaction
from django.db.models import Count, Q, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render

from accounts.models import Ward
from accounts.permissions import has_perm, user_perms
from accounts.wards import sorted_wards

from .budget import TOL, ZERO, D, check_project_budget, distributed, head_room, rs, usage
from .models import Budget_Type, BudgetAmendment, BudgetHead, FiscalYear, Project


def _user_wards(user):
    from usercommittee.views_workflow import _user_wards as uw

    return uw(user)


def _money(raw):
    try:
        v = Decimal(str(raw or "").replace(",", "").strip().translate(str.maketrans("०१२३४५६७८९", "0123456789")))
    except (InvalidOperation, ValueError):
        return None
    return v


def _fy(request):
    """छानिएको आ.व. — नभए पछिल्लो."""
    v = request.GET.get("fiscalyear") or request.POST.get("fiscalyear") or ""
    fy = FiscalYear.objects.filter(id=v).first() if str(v).isdigit() else None
    return fy or FiscalYear.objects.order_by("-name").first()


# ---------------------------------------------------------------- बजेट शीर्षक (सेटिङ)
@login_required
def budget_heads(request):
    if not has_perm(request.user, "budget.manage"):
        return redirect("unauthorized")
    fy = _fy(request)
    if request.method == "POST":
        action = request.POST.get("action", "")
        code = (request.POST.get("code") or "").strip()
        name = (request.POST.get("name") or "").strip()
        amount = _money(request.POST.get("allocated"))
        bt = request.POST.get("budget_type") or ""
        budget_type = Budget_Type.objects.filter(id=bt).first() if bt.isdigit() else None
        if action in ("add", "edit"):
            head = get_object_or_404(BudgetHead, id=request.POST.get("id")) if action == "edit" else None
            dup = BudgetHead.objects.filter(fiscalyear=fy, code=code).exclude(id=getattr(head, "id", None))
            if not name:
                messages.error(request, "शीर्षकको नाम लेख्नुहोस्।")
            elif amount is None or amount < 0:
                messages.error(request, "स्वीकृत रकम अङ्कमा लेख्नुहोस्।")
            elif code and dup.exists():
                messages.error(request, f"उपशीर्षक नं. “{code}” {fy.name} मा पहिले नै छ।")
            elif head is not None and amount + TOL < distributed(head):
                messages.error(request, f"यस शीर्षकका योजनामा {rs(distributed(head))} विनियोजन भइसकेको छ — "
                                        "स्वीकृत रकम त्योभन्दा कम गर्न मिल्दैन।")
            else:
                if head is None:
                    head = BudgetHead(fiscalyear=fy)
                head.code, head.name, head.allocated, head.budget_type = code, name, amount, budget_type
                head.save()
                messages.success(request, f"बजेट शीर्षक “{head}” ({rs(amount)}) सुरक्षित भयो।")
        elif action == "toggle":
            head = get_object_or_404(BudgetHead, id=request.POST.get("id"))
            head.is_active = not head.is_active
            head.save(update_fields=["is_active"])
            messages.success(request, f"“{head}” {'सक्रिय' if head.is_active else 'निष्क्रिय'} गरियो।")
        elif action == "delete":
            head = get_object_or_404(BudgetHead, id=request.POST.get("id"))
            if head.projects.exists():
                messages.error(request, f"“{head}” मा योजना राखिसकेकाले हटाउन मिल्दैन — निष्क्रिय गर्नुहोस्।")
            else:
                head.delete()
                messages.success(request, f"“{head}” हटाइयो।")
        return redirect(f"{request.path}?fiscalyear={fy.id}" if fy else request.path)

    heads = list(BudgetHead.objects.filter(fiscalyear=fy).select_related("budget_type")
                 .annotate(n=Count("projects"), dist=Sum("projects__budget")))
    for h in heads:
        h.dist = D(h.dist)
        h.room = D(h.allocated) - h.dist
    edit_id = request.GET.get("edit", "")
    return render(request, "projects/budget_heads.html", {
        "fy": fy, "fiscal_years": FiscalYear.objects.order_by("-name"), "heads": heads,
        "budget_types": Budget_Type.objects.all(), "edit_id": int(edit_id) if edit_id.isdigit() else None,
        "total": sum((D(h.allocated) for h in heads), ZERO), "total_dist": sum((h.dist for h in heads), ZERO),
        "unassigned": Project.objects.filter(fiscalyear=fy, budget_head__isnull=True).count() if fy else 0,
    })


# ---------------------------------------------------------------- थप / कटौती / रकमान्तर
@login_required
def amendments(request):
    if not user_perms(request.user) & {"budget.view", "budget.manage"}:
        return redirect("unauthorized")
    fy = _fy(request)
    qs = (BudgetAmendment.objects.select_related("from_project__ward", "to_project__ward", "created_by")
          .filter(Q(from_project__fiscalyear=fy) | Q(to_project__fiscalyear=fy)).distinct())
    wards = _user_wards(request.user)
    if wards is not None:
        qs = qs.filter(Q(from_project__ward_id__in=wards) | Q(to_project__ward_id__in=wards))
    return render(request, "projects/budget_amendments.html", {
        "fy": fy, "fiscal_years": FiscalYear.objects.order_by("-name"), "rows": qs,
        "can_manage": has_perm(request.user, "budget.manage"),
    })


@login_required
def amendment_new(request):
    if not has_perm(request.user, "budget.manage"):
        return redirect("unauthorized")
    from biddings.views_workflow import to_ad, to_bs

    fy = _fy(request)
    projects = list(Project.objects.filter(fiscalyear=fy).select_related("ward", "budget_head").order_by("ward_id", "name"))
    use = usage([p.id for p in projects])
    for p in projects:
        u = use.get(p.id, {})
        p.spent = u.get("spent", ZERO)
        p.free = D(p.budget) - max(p.spent, u.get("committed", ZERO))
    errors = []
    data = request.POST if request.method == "POST" else request.GET
    kind = data.get("kind") if data.get("kind") in dict(BudgetAmendment.KINDS) else "transfer"
    if request.method == "POST":
        pid = {p.id: p for p in projects}
        src = pid.get(int(data["from_project"])) if str(data.get("from_project", "")).isdigit() else None
        dst = pid.get(int(data["to_project"])) if str(data.get("to_project", "")).isdigit() else None
        if kind == "increase":
            src = None          # लुकेको फिल्डको मान बेवास्ता
        if kind == "decrease":
            dst = None
        amount = _money(data.get("amount"))
        day = to_ad(data.get("decision_date"))
        reason = (data.get("reason") or "").strip()
        if kind in ("decrease", "transfer") and src is None:
            errors.append("रकम घटाउने योजना छान्नुहोस्।")
        if kind in ("increase", "transfer") and dst is None:
            errors.append("रकम थप्ने योजना छान्नुहोस्।")
        if kind == "transfer" and src is not None and src == dst:
            errors.append("एउटै योजनामा रकमान्तर हुँदैन।")
        if amount is None or amount <= 0:
            errors.append("रकम ० भन्दा बढी अङ्कमा लेख्नुहोस्।")
        if day is None:
            errors.append("निर्णय मिति भर्नुहोस् (वि.सं. YYYY-MM-DD)।")
        if not reason:
            errors.append("कारण लेख्नुहोस्।")
        if not errors:
            if src is not None:
                errors += check_project_budget(src, D(src.budget) - amount, None)
            if dst is not None and dst.budget_head_id:
                room = head_room(dst.budget_head, [dst.id])
                if src is not None and src.budget_head_id == dst.budget_head_id:
                    room += amount       # एउटै शीर्षकभित्रको रकमान्तर
                if D(dst.budget) + amount > room + TOL:
                    errors.append(f"“{dst.budget_head}” शीर्षकमा बाँकी {rs(head_room(dst.budget_head, [dst.id]) - D(dst.budget))} "
                                  "मात्र छ — पहिले शीर्षकको स्वीकृत रकम बढाउनुहोस् वा त्यही शीर्षकको योजनाबाट रकमान्तर गर्नुहोस्।")
        if not errors:
            with transaction.atomic():
                a = BudgetAmendment(kind=kind, amount=amount, decision_date=day, reason=reason,
                                    decision_by=(data.get("decision_by") or "").strip(),
                                    decision_no=(data.get("decision_no") or "").strip(),
                                    document=request.FILES.get("document"), created_by=request.user)
                if src is not None:
                    src = Project.objects.select_for_update().get(id=src.id)
                    a.from_project, a.from_before = src, D(src.budget)
                    a.from_after = a.from_before - amount
                    src.budget = float(a.from_after)
                    src.save(update_fields=["budget"])
                if dst is not None:
                    dst = Project.objects.select_for_update().get(id=dst.id)
                    a.to_project, a.to_before = dst, D(dst.budget)
                    a.to_after = a.to_before + amount
                    dst.budget = float(a.to_after)
                    dst.save(update_fields=["budget"])
                a.save()
            messages.success(request, f"{a.get_kind_display()} {rs(amount)} दर्ता भयो र विनियोजित रकम मिलाइयो।")
            return redirect(f"/Projects/Budget/Amendments/?fiscalyear={fy.id}")
    return render(request, "projects/budget_amendment_form.html", {
        "fy": fy, "fiscal_years": FiscalYear.objects.order_by("-name"), "projects": projects, "kind": kind,
        "kinds": BudgetAmendment.KINDS, "errors": errors, "d": data,
        "today_bs": to_bs(__import__("datetime").date.today()),
    })


# ---------------------------------------------------------------- बजेट स्थिति प्रतिवेदन
def _status_rows(request):
    fy = _fy(request)
    f = {k: (request.GET.get(k) or "").strip() for k in ("ward", "budget_type", "q", "flag")}
    projects = Project.objects.filter(fiscalyear=fy).select_related("ward", "budget_head", "procurement_type", "budget_type")
    wards = _user_wards(request.user)
    if wards is not None:
        projects = projects.filter(ward_id__in=wards)
    if f["ward"].isdigit():
        projects = projects.filter(ward_id=f["ward"])
    if f["budget_type"].isdigit():
        projects = projects.filter(Q(budget_head__budget_type_id=f["budget_type"]) |
                                   Q(budget_head__isnull=True, budget_type_id=f["budget_type"]))
    if f["q"]:
        projects = projects.filter(Q(name__icontains=f["q"]) | Q(budget_head__name__icontains=f["q"]) |
                                   Q(budget_head__code__icontains=f["q"]))
    projects = list(projects.order_by("ward_id", "name"))
    use = usage([p.id for p in projects])
    rows = []
    for p in projects:
        u = use.get(p.id, {})
        budget, sp, com = D(p.budget), u.get("spent", ZERO), u.get("committed", ZERO)
        r = u.get("row")
        flag = "over" if sp > budget + TOL or com > budget + TOL else ("near" if budget and sp >= budget * Decimal("0.9") else "")
        rows.append({"p": p, "budget": budget, "spent": sp, "committed": com, "balance": budget - sp,
                     "free": budget - max(sp, com), "pct": round(sp * 100 / budget, 1) if budget else 0,
                     "flag": flag, "kind_label": r["kind_label"] if r else (p.procurement_type.name if p.procurement_type else ""),
                     "status_label": r["status_label"] if r else ""})
    if f["flag"] in ("over", "near"):
        rows = [r for r in rows if r["flag"] == f["flag"]]

    groups = OrderedDict()
    heads = BudgetHead.objects.filter(fiscalyear=fy).select_related("budget_type") if fy else []
    for h in heads:
        groups[h.id] = {"head": h, "rows": []}
    groups[None] = {"head": None, "rows": []}
    for r in rows:
        groups.setdefault(r["p"].budget_head_id, {"head": r["p"].budget_head, "rows": []})["rows"].append(r)
    out = []
    filtered = bool(f["ward"] or f["q"] or f["flag"] or wards is not None)
    for key, g in groups.items():
        h, rs_ = g["head"], g["rows"]
        if f["budget_type"].isdigit() and h is not None and str(h.budget_type_id) != f["budget_type"]:
            continue
        if not rs_ and (h is None or filtered):
            continue
        t = {k: sum((r[k] for r in rs_), ZERO) for k in ("budget", "spent", "committed", "balance")}
        g.update(t)
        g["allocated"] = D(h.allocated) if h else None
        g["unassigned"] = (D(h.allocated) - distributed(h)) if h else None
        g["pct"] = round(t["spent"] * 100 / t["budget"], 1) if t["budget"] else 0
        out.append(g)
    total = {k: sum((g[k] for g in out), ZERO) for k in ("budget", "spent", "committed", "balance")}
    total["allocated"] = sum((g["allocated"] for g in out if g["allocated"] is not None), ZERO)
    total["pct"] = round(total["spent"] * 100 / total["budget"], 1) if total["budget"] else 0
    total["over"] = sum(1 for r in rows if r["flag"] == "over")
    total["near"] = sum(1 for r in rows if r["flag"] == "near")
    total["count"] = len(rows)
    keep = request.GET.copy()
    keep.pop("export", None)
    if fy and "fiscalyear" not in keep:
        keep["fiscalyear"] = str(fy.id)
    return fy, f, out, total, keep.urlencode(), wards


@login_required
def budget_status(request):
    if not user_perms(request.user) & {"budget.view", "budget.manage"}:
        return redirect("unauthorized")
    fy, f, groups, total, qs, wards = _status_rows(request)
    title = "बजेट स्थिति (बजेट नियन्त्रण)"
    if request.GET.get("export") == "xlsx":
        return _xlsx(fy, groups, total)
    desc = [f"आ.व. {fy.name}" if fy else ""]
    if f["ward"].isdigit():
        w = Ward.objects.filter(id=f["ward"]).first()
        desc.append(f"वडा नं. {w.name}" if w else "")
    if f["flag"]:
        desc.append({"over": "विनियोजनभन्दा बढी", "near": "९०% भन्दा बढी खर्च"}[f["flag"]])
    return render(request, "projects/budget_status.html", {
        "title": title, "fy": fy, "f": f, "groups": groups, "total": total, "qs": qs,
        "filter_desc": " · ".join(x for x in desc if x),
        "fiscal_years": FiscalYear.objects.order_by("-name"), "budget_types": Budget_Type.objects.all(),
        "all_wards": sorted_wards(Ward.objects.filter(id__in=wards) if wards is not None else None),
        "can_manage": has_perm(request.user, "budget.manage"),
    })


def _xlsx(fy, groups, total):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "शीर्षकअनुसार"
    head = ["उपशीर्षक नं.", "बजेट शीर्षक", "स्रोत", "स्वीकृत रकम", "योजनामा विनियोजित", "योजना नछुट्याएको",
            "प्रतिबद्धता", "खर्च", "बाँकी (विनियोजित − खर्च)", "खर्च %"]
    ws.append(head)
    for g in groups:
        h = g["head"]
        ws.append([h.code if h else "", str(h.name) if h else "शीर्षक नतोकिएका योजना",
                   (h.budget_type.name if h and h.budget_type else ""),
                   float(g["allocated"]) if g["allocated"] is not None else None, float(g["budget"]),
                   float(g["unassigned"]) if g["unassigned"] is not None else None,
                   float(g["committed"]), float(g["spent"]), float(g["balance"]), float(g["pct"])])
    ws.append(["", "जम्मा", "", float(total["allocated"]), float(total["budget"]), None, float(total["committed"]),
               float(total["spent"]), float(total["balance"]), float(total["pct"])])
    ws2 = wb.create_sheet("योजनाअनुसार")
    ws2.append(["क्र.सं.", "बजेट शीर्षक", "योजनाको नाम", "वडा", "खरिद विधि", "स्थिति", "विनियोजित", "प्रतिबद्धता",
                "खर्च", "बाँकी", "खर्च %", "कैफियत"])
    i = 0
    for g in groups:
        for r in g["rows"]:
            i += 1
            ws2.append([i, str(g["head"]) if g["head"] else "", r["p"].name, getattr(r["p"].ward, "name", ""),
                        r["kind_label"], r["status_label"], float(r["budget"]), float(r["committed"]),
                        float(r["spent"]), float(r["balance"]), float(r["pct"]),
                        {"over": "विनियोजनभन्दा बढी", "near": "९०%+ खर्च"}.get(r["flag"], "")])
    for sheet, widths in ((ws, [12, 36, 16, 16, 16, 16, 16, 16, 18, 8]), (ws2, [6, 30, 36, 6, 14, 14, 15, 15, 15, 15, 8, 16])):
        for c in sheet[1]:
            c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E8EEFB")
        for n, w in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=n).column_letter].width = w
        sheet.freeze_panes = "A2"
    resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = 'attachment; filename="budget-status.xlsx"'
    wb.save(resp)
    return resp
