"""कर तथा धरौटी कट्टी विवरण र धरौटी फिर्ता (सिफारिस अपलोड → टिप्पणी → फिर्ता दर्ता)."""
from collections import OrderedDict
from datetime import date

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone

from accounts.models import Ward
from accounts.permissions import has_perm, user_perms
from accounts.wards import sorted_wards
from usercommittee.workflow import bs_date

from .deductions import (KIND_PTYPE, KINDS, STATUS, ZERO, D, all_lines, kind_of, labels, projects_qs, retention_item,
                         retention_items, retention_months)
from .models import FiscalYear, Project, RetentionRelease

UPLOAD_TYPES = (".pdf", ".jpg", ".jpeg", ".png")
UPLOAD_MAX_MB = 10


def _user_wards(user):
    from usercommittee.views_workflow import _user_wards as uw

    return uw(user)


def _base(request, default_latest=True):
    """(project_ids, ctx) — आ.व., वडा, खरिद विधि फिल्टर; वडा तहका प्रयोगकर्ताले आफ्नो वडा मात्र."""
    qs = Project.objects.filter(procurement_type_id__in=KIND_PTYPE.values())
    my_wards = _user_wards(request.user)
    if my_wards is not None:
        qs = qs.filter(ward_id__in=my_wards)
    fys = list(FiscalYear.objects.all().order_by("-id"))
    fy = request.GET.get("fiscalyear")
    if fy is None:
        fy = str(fys[0].id) if fys and default_latest else "all"
    ward = request.GET.get("ward", "")
    kind = request.GET.get("kind", "")
    kind = kind if kind in KINDS else ""
    if fy.isdigit():
        qs = qs.filter(fiscalyear_id=fy)
    if ward.isdigit():
        qs = qs.filter(ward_id=ward)
    if kind:
        qs = qs.filter(procurement_type_id=KIND_PTYPE[kind])
    all_wards = sorted_wards(Ward.objects.filter(id__in=my_wards) if my_wards is not None else None)
    fy_name = next((x.name for x in fys if str(x.id) == fy), "")
    desc = [KINDS[kind] if kind else "उपभोक्ता समिति, बोलपत्र र दरभाउपत्र",
            f"आ.व. {fy_name}" if fy_name else "सबै आर्थिक वर्ष",
            next((f"वडा नं. {w.name}" for w in all_wards if str(w.id) == ward), "सबै वडा")]
    keep = request.GET.copy()
    keep.pop("export", None)
    keep["fiscalyear"] = fy
    return list(qs.values_list("id", flat=True)), {
        "fiscal_years": fys, "all_wards": all_wards, "kinds": KINDS.items(),
        "f": {"fiscalyear": fy, "ward": ward, "kind": kind, "code": request.GET.get("code", ""),
              "status": request.GET.get("status", ""), "q": request.GET.get("q", "").strip()},
        "filter_desc": " · ".join(desc), "qs": keep.urlencode(),
    }


# ---------------------------------------------------------------- कर तथा धरौटी कट्टी विवरण
@login_required
def deduction_register(request):
    if not has_perm(request.user, "report.uc"):
        return redirect("unauthorized")
    ids, ctx = _base(request)
    code = ctx["f"]["code"]
    lines = all_lines(ids, ctx["f"]["kind"], code)
    q = ctx["f"]["q"]
    if q:
        lines = [ln for ln in lines if q in ln["party"] or q in ln["p"].name]
    for ln in lines:
        ln["date_bs"] = bs_date(ln["date"])
    lab = labels()
    heads = OrderedDict()
    parties = OrderedDict()
    for ln in lines:
        h = heads.setdefault(ln["code"], {"code": ln["code"], "label": ln["label"], "total": ZERO, "count": 0,
                                          "retention": ln["retention"]})
        h["total"] += ln["amount"]
        h["count"] += 1
        key = (ln["kind"], ln["party"])
        pt = parties.setdefault(key, {"party": ln["party"] or "—", "kind_label": ln["kind_label"], "total": ZERO,
                                      "by_code": {}, "projects": set()})
        pt["total"] += ln["amount"]
        pt["by_code"][ln["code"]] = pt["by_code"].get(ln["code"], ZERO) + ln["amount"]
        pt["projects"].add(ln["p"].id)
    head_list = sorted(heads.values(), key=lambda h: (h["retention"], h["label"]))
    party_list = sorted(parties.values(), key=lambda x: -x["total"])
    for pt in party_list:
        pt["values"] = [pt["by_code"].get(h["code"]) for h in head_list]
        pt["n_projects"] = len(pt["projects"])
    total = sum((ln["amount"] for ln in lines), ZERO)
    if code:
        ctx["filter_desc"] += " · " + lab.get(code, code)
    if request.GET.get("export") == "xlsx":
        return _register_xlsx(lines, head_list, party_list, ctx)
    all_codes = sorted({(c, lab.get(c, c)) for c in lab} | {(h["code"], h["label"]) for h in head_list},
                       key=lambda x: x[1])
    return render(request, "projects/deduction_register.html", ctx | {
        "lines": lines, "heads": head_list, "parties": party_list, "total": total, "codes": all_codes,
        "retention_total": sum((h["total"] for h in head_list if h["retention"]), ZERO),
    })


def _register_xlsx(lines, heads, parties, ctx):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    fill = PatternFill("solid", fgColor="E8EEFB")

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), fill
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        ws.freeze_panes = "A2"

    ws = wb.active
    ws.title = "कट्टी विवरण"
    sheet(ws, ["मिति (वि.सं.)", "योजना", "वडा", "खरिद विधि", "समिति / निर्माण व्यवसायी / फर्म", "स्रोत", "शीर्षक", "रकम"],
          [[bs_date(ln["date"]), ln["p"].name, getattr(ln["p"].ward, "name", ""), ln["kind_label"], ln["party"],
            ln["source"], ln["label"], float(ln["amount"])] for ln in lines],
          [13, 36, 6, 15, 30, 16, 26, 14])
    sheet(wb.create_sheet("शीर्षकअनुसार"), ["शीर्षक", "पटक", "जम्मा"],
          [[h["label"], h["count"], float(h["total"])] for h in heads], [30, 8, 16])
    sheet(wb.create_sheet("पक्षअनुसार"), ["समिति / निर्माण व्यवसायी / फर्म", "खरिद विधि", "योजना"] + [h["label"] for h in heads]
          + ["जम्मा"],
          [[p["party"], p["kind_label"], p["n_projects"]] + [float(v) if v else None for v in p["values"]]
           + [float(p["total"])] for p in parties],
          [30, 15, 8] + [16] * len(heads) + [16])
    resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = 'attachment; filename="kar-dharauti-katti.xlsx"'
    wb.save(resp)
    return resp


# ---------------------------------------------------------------- धरौटी फिर्ता
def _can_see_retention(user):
    return bool(user_perms(user) & {"retention.recommend", "retention.manage", "report.uc"})


@login_required
def retention_list(request):
    if not _can_see_retention(request.user):
        return redirect("unauthorized")
    ids, ctx = _base(request, default_latest=False)
    items = retention_items(ids, ctx["f"]["kind"])
    q = ctx["f"]["q"]
    if q:
        items = [it for it in items if q in it["party"] or q in it["p"].name]
    counts = {k: 0 for k in STATUS}
    for it in items:
        counts[it["status"]] += 1
    status = ctx["f"]["status"]
    if status == "pending":
        shown = [it for it in items if it["status"] in ("due", "recommended", "tippani")]
    elif status in STATUS:
        shown = [it for it in items if it["status"] == status]
    else:
        shown = items
    order = {"due": 0, "recommended": 1, "tippani": 2, "held": 3, "running": 4, "released": 5}
    shown.sort(key=lambda it: (order[it["status"]], it["due"] or date.max))
    for it in shown:
        it["completed_bs"], it["due_bs"] = bs_date(it["completed_on"]), bs_date(it["due"])
        rel = it["release"]
        it["rec_bs"] = bs_date(rel.rec_date) if rel else ""
        it["released_bs"] = bs_date(rel.released_on) if rel and rel.released_on else ""
        it["tippani_bs"] = bs_date(timezone.localdate(rel.tippani_printed_at)) if rel and rel.tippani_printed_at else ""
    held_total = sum((it["held"] for it in items if it["status"] != "released"), ZERO)
    released_total = sum((D(it["release"].released_amount) for it in items if it["status"] == "released"), ZERO)
    P = user_perms(request.user)
    return render(request, "projects/retention_list.html", ctx | {
        "items": shown, "counts": counts, "statuses": STATUS, "months": retention_months(),
        "held_total": held_total, "released_total": released_total,
        "can_recommend": "retention.recommend" in P, "can_manage": "retention.manage" in P,
        "open_id": int(request.GET["open"]) if request.GET.get("open", "").isdigit() else None,
    })


def _item_or_redirect(request, project_id):
    project = get_object_or_404(Project, id=project_id)
    wards = _user_wards(request.user)
    if wards is not None and project.ward_id not in wards:
        return None, redirect("unauthorized")
    kind = kind_of(project)
    if kind is None:
        return None, redirect("retention_list")
    p = projects_qs(kind).select_related("retention_release").filter(id=project.id).first()
    item = retention_item(p, kind) if p else None
    if item is None:
        messages.error(request, f"“{project.name}” मा धरौटी कट्टी भएको छैन।")
        return None, redirect("retention_list")
    return item, None


def _back(project_id):
    return reverse("retention_list") + f"?fiscalyear=all&open={project_id}"


def _check_file(f, errors, label):
    if not f:
        errors.append(f"{label} छान्नुहोस्।")
    elif not f.name.lower().endswith(UPLOAD_TYPES):
        errors.append(f"{label}: PDF, JPG वा PNG मात्र।")
    elif f.size > UPLOAD_MAX_MB * 1024 * 1024:
        errors.append(f"{label}: फाइल {UPLOAD_MAX_MB} MB भन्दा सानो हुनुपर्छ।")


@login_required
def retention_recommend(request, project_id):
    """इन्जिनियर / वडाको सिफारिस पत्र अपलोड (फिर्ता अघि सम्म सच्याउन मिल्ने)."""
    if not has_perm(request.user, "retention.recommend"):
        messages.error(request, "धरौटी फिर्ता सिफारिस अपलोड गर्ने अनुमति छैन।")
        return redirect("retention_list")
    item, r = _item_or_redirect(request, project_id)
    if r:
        return r
    rel = item["release"]
    if rel is not None and rel.released_on:
        messages.error(request, "धरौटी फिर्ता भइसकेको छ।")
        return redirect(_back(project_id))
    if item["status"] == "running":
        messages.error(request, "अन्तिम भुक्तानी नभई धरौटी फिर्ता सिफारिस गर्न मिल्दैन।")
        return redirect(_back(project_id))
    from biddings.views_workflow import Form, to_bs

    errors = []
    if request.method == "POST":
        f = Form(request)
        rec_date = f.day("rec_date", "सिफारिस मिति")
        rec_by = request.POST.get("rec_by") if request.POST.get("rec_by") in dict(RetentionRelease.REC_BY) else "engineer"
        doc = request.FILES.get("rec_document")
        errors = f.errors
        if doc or rel is None:
            _check_file(doc, errors, "सिफारिस पत्रको स्क्यान")
        if rec_date and item["completed_on"] and rec_date < item["completed_on"]:
            errors.append("सिफारिस मिति अन्तिम भुक्तानी मितिभन्दा अघि हुन सक्दैन।")
        if not errors:
            obj = rel or RetentionRelease(project=item["p"])
            obj.held_amount = item["held"]
            obj.rec_by, obj.rec_date = rec_by, rec_date
            obj.rec_number = request.POST.get("rec_number", "").strip()
            obj.rec_remarks = request.POST.get("rec_remarks", "").strip()
            obj.rec_uploaded_by = request.user
            if doc:
                obj.rec_document = doc
            obj.tippani_printed_at = None if doc and rel is not None else obj.tippani_printed_at
            obj.save()
            if item["due"] and item["due"] > date.today():
                messages.warning(request, f"धरौटी अवधि {bs_date(item['due'])} मा मात्र सकिन्छ — म्यादअघि फिर्ता गर्दा कारण खुलाउनुपर्छ।")
            messages.success(request, f"“{item['p'].name}” को धरौटी फिर्ता सिफारिस अपलोड भयो। अब टिप्पणी तयार गर्नुहोस्।")
            return redirect(_back(project_id))
    post = request.POST if request.method == "POST" else {}
    return render(request, "projects/retention_form.html", {
        "item": item, "mode": "recommend", "errors": errors, "rec_by_choices": RetentionRelease.REC_BY,
        "v": {"rec_date": post.get("rec_date", to_bs(rel.rec_date) if rel else to_bs(date.today())),
              "rec_number": post.get("rec_number", rel.rec_number if rel else ""),
              "rec_by": post.get("rec_by", rel.rec_by if rel else ("ward" if request.user.user_type in ("6", "7") else "engineer")),
              "rec_remarks": post.get("rec_remarks", rel.rec_remarks if rel else "")},
        "back": _back(project_id), "due_bs": bs_date(item["due"]), "completed_bs": bs_date(item["completed_on"]),
    })


@login_required
def retention_tippani(request, project_id):
    """धरौटी फिर्ता टिप्पणी (सिफारिस आएपछि मात्र)."""
    if not has_perm(request.user, "retention.manage"):
        messages.error(request, "धरौटी फिर्ता टिप्पणी छाप्ने अनुमति छैन।")
        return redirect("retention_list")
    item, r = _item_or_redirect(request, project_id)
    if r:
        return r
    rel = item["release"]
    if rel is None:
        messages.error(request, f"“{item['p'].name}”: पहिले इन्जिनियर / वडाको सिफारिस अपलोड गर्नुहोस्।")
        return redirect(_back(project_id))
    if not rel.tippani_printed_at:
        rel.tippani_printed_at = timezone.now()
        rel.save(update_fields=["tippani_printed_at"])
    p = item["p"]
    return render(request, "projects/print_retention_tippani.html", {
        "item": item, "rel": rel, "project": p, "today_bs": bs_date(date.today()),
        "completed_bs": bs_date(item["completed_on"]), "due_bs": bs_date(item["due"]), "rec_bs": bs_date(rel.rec_date),
        "months": retention_months(), "early": bool(item["due"] and item["due"] > date.today()),
        "lines": [ln | {"date_bs": bs_date(ln["date"])} for ln in item["lines"]],
        "back": _back(project_id),
    })


@login_required
def retention_release(request, project_id):
    """फिर्ता भएको दर्ता: मिति, रकम, भौचर, सदर भएको टिप्पणी (स्क्यान)."""
    if not has_perm(request.user, "retention.manage"):
        messages.error(request, "धरौटी फिर्ता दर्ता गर्ने अनुमति छैन।")
        return redirect("retention_list")
    item, r = _item_or_redirect(request, project_id)
    if r:
        return r
    rel = item["release"]
    if rel is None or not rel.tippani_printed_at:
        messages.error(request, f"“{item['p'].name}”: सिफारिस र टिप्पणी पछि मात्र फिर्ता दर्ता गर्न मिल्छ।")
        return redirect(_back(project_id))
    if rel.released_on:
        messages.error(request, "धरौटी फिर्ता भइसकेको छ।")
        return redirect(_back(project_id))
    from biddings.views_workflow import Form, plain, to_bs

    early = bool(item["due"] and item["due"] > date.today())
    errors = []
    if request.method == "POST":
        f = Form(request)
        on = f.day("released_on", "फिर्ता मिति")
        amount = f.amount("released_amount", "फिर्ता रकम")
        doc = request.FILES.get("approved_document")
        reason = request.POST.get("early_reason", "").strip()
        errors = f.errors
        _check_file(doc, errors, "सदर भएको टिप्पणीको स्क्यान")
        if amount is not None and amount > item["held"]:
            errors.append(f"फिर्ता रकम रोक्का धरौटी (रु. {item['held']:,.2f}) भन्दा बढी हुन सक्दैन।")
        if on and on < rel.rec_date:
            errors.append("फिर्ता मिति सिफारिस मितिभन्दा अघि हुन सक्दैन।")
        if early and not reason:
            errors.append(f"धरौटी अवधि {bs_date(item['due'])} मा सकिन्छ — म्यादअघि फिर्ताको कारण लेख्नुहोस्।")
        if not errors:
            rel.released_on, rel.released_amount = on, amount
            rel.voucher_number = request.POST.get("voucher_number", "").strip()
            rel.approved_document, rel.early_reason, rel.released_by = doc, reason, request.user
            rel.held_amount = item["held"]
            rel.save()
            messages.success(request, f"“{item['p'].name}” को धरौटी रु. {amount:,.2f} फिर्ता दर्ता भयो।")
            return redirect(_back(project_id))
    post = request.POST if request.method == "POST" else {}
    return render(request, "projects/retention_form.html", {
        "item": item, "mode": "release", "errors": errors, "early": early,
        "v": {"released_on": post.get("released_on", to_bs(date.today())),
              "released_amount": post.get("released_amount", plain(item["held"])),
              "voucher_number": post.get("voucher_number", ""), "early_reason": post.get("early_reason", "")},
        "back": _back(project_id), "due_bs": bs_date(item["due"]), "completed_bs": bs_date(item["completed_on"]),
    })
