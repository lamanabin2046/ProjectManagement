"""बैंक ग्यारेन्टी: कार्यसम्पादन जमानत (Performance Bank Guarantee) र अग्रिम भुक्तानी जमानत (Advance Payment Guarantee).

PG  — बोलपत्र ठेक्का (PerformanceBondIssued; म्याद थप भए पछिल्लो म्याद)
APG — बोलपत्र (AdvanceBankGuarantee) र दरभाउपत्र (QuationAdvance)

स्थिति: active (३० दिनभन्दा बढी बाँकी) · expiring (३० दिनभित्र) · expired (म्याद सकियो) · closed (फिर्ता/समायोजन भइसकेको)
"""
from datetime import date
from decimal import Decimal

from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import redirect, render

from accounts.models import Ward
from accounts.permissions import user_perms
from accounts.wards import sorted_wards
from usercommittee.workflow import bs_date

from .models import FiscalYear

WARN_DAYS = 30
STATUS = {"active": "बहाल", "expiring": "म्याद सकिन लागेको", "expired": "म्याद सकिएको", "closed": "फिर्ता / समायोजन भइसकेको"}
ZERO = Decimal("0")


def _user_wards(user):
    from usercommittee.views_workflow import _user_wards as uw

    return uw(user)


def _status(expiry, closed):
    if closed:
        return "closed"
    days = (expiry - date.today()).days if expiry else None
    if days is None or days < 0:
        return "expired"
    return "expiring" if days <= WARN_DAYS else "active"


def _row(project, kind, party, bank, number, amount, issue, expiry, closed, extended=False, against=None, links=None):
    st = _status(expiry, closed)
    return {"p": project, "kind": kind, "kind_label": {"bid": "बोलपत्र ठेक्का", "quot": "दरभाउपत्र खरिद"}[kind],
            "party": party or "", "bank": bank, "number": number, "amount": Decimal(str(amount or 0)),
            "issue": issue, "expiry": expiry, "issue_bs": bs_date(issue), "expiry_bs": bs_date(expiry),
            "days": (expiry - date.today()).days if expiry else None, "status": st, "status_label": STATUS[st],
            "extended": extended, "against": against, "links": links or []}


def pg_rows():
    from biddings.models import PerformanceBondIssued

    rows = []
    qs = (PerformanceBondIssued.objects
          .select_related("contract__sealed_bid__project__ward", "contract__sealed_bid__project__fiscalyear",
                          "contract__sealed_bid__project__retention_release")
          .prefetch_related("extensions"))
    for b in qs:
        bid = b.contract.sealed_bid
        p = bid.project
        rel = getattr(p, "retention_release", None)
        closed = bool(rel and rel.released_on)          # धरौटी फिर्ता भएपछि PG पनि फिर्ता हुन्छ
        exts = list(b.extensions.all())
        rows.append(_row(p, "bid", bid.bidder_name_nepali, b.bank_name, b.bond_number, b.issued_amount, b.issue_date,
                         b.current_expiry_date, closed, extended=bool(exts),
                         against=bid.bidding_amount,
                         links=[("bond", b.id)]))
    return rows


def apg_rows():
    from biddings.models import AdvanceBankGuarantee
    from quatation.models import QuationAdvance
    from quatation.workflow import final_of, lowest

    rows = []
    for g in AdvanceBankGuarantee.objects.select_related(
            "advance_payment__contract__sealed_bid__project__ward", "advance_payment__contract__sealed_bid__project__fiscalyear",
            "advance_payment__contract__final_payment"):
        ap = g.advance_payment
        contract = ap.contract
        bid = contract.sealed_bid
        closed = getattr(contract, "final_payment", None) is not None   # अन्तिम भुक्तानीमा अग्रिम समायोजन
        rows.append(_row(bid.project, "bid", bid.bidder_name_nepali, g.bank_name, g.guarantee_number,
                         g.guaranteed_amount, g.issue_date, g.expiry_date, closed, against=ap.approved_amount,
                         links=[("bid_advance", ap.id)]))
    for a in QuationAdvance.objects.select_related("quation__project__ward", "quation__project__fiscalyear",
                                                   "quation__approval").prefetch_related("quation__quationfinal_set"):
        q = a.quation
        appr = getattr(q, "approval", None)
        firm = appr.firm_name if appr and appr.firm_name else lowest(q)[0]
        rows.append(_row(q.project, "quot", firm, a.bank_name, a.guarantee_number, a.guaranteed_amount, a.issue_date,
                         a.expiry_date, final_of(q) is not None, against=a.approved_amount,
                         links=[("quot_advance", a.id)]))
    return rows


def _filtered(request, rows):
    wards = _user_wards(request.user)
    if wards is not None:
        rows = [r for r in rows if r["p"].ward_id in wards]
    f = {k: request.GET.get(k, "") for k in ("status", "kind", "fiscalyear", "ward", "q")}
    f["q"] = f["q"].strip()
    if f["fiscalyear"].isdigit():
        rows = [r for r in rows if str(r["p"].fiscalyear_id) == f["fiscalyear"]]
    if f["ward"].isdigit():
        rows = [r for r in rows if str(r["p"].ward_id) == f["ward"]]
    if f["kind"] in ("bid", "quot"):
        rows = [r for r in rows if r["kind"] == f["kind"]]
    if f["q"]:
        q = f["q"].lower()
        rows = [r for r in rows if q in r["p"].name.lower() or q in (r["party"] or "").lower()
                or q in (r["bank"] or "").lower() or q in (r["number"] or "").lower()]
    counts = {k: sum(1 for r in rows if r["status"] == k) for k in STATUS}   # स्थिति बाहेकका फिल्टरपछि
    if f["status"] == "alert":
        rows = [r for r in rows if r["status"] in ("expiring", "expired")]
    elif f["status"] in STATUS:
        rows = [r for r in rows if r["status"] == f["status"]]
    elif not f["status"]:
        rows = [r for r in rows if r["status"] != "closed"]       # पूर्वनिर्धारित: बहाल मात्र (फिर्ता भएका बाहेक)
    order = {"expired": 0, "expiring": 1, "active": 2, "closed": 3}
    rows.sort(key=lambda r: (order[r["status"]], r["expiry"] or date.max))
    keep = request.GET.copy()
    keep.pop("export", None)
    return rows, counts, f, keep.urlencode(), wards


def _page(request, which):
    P = user_perms(request.user)
    if not P & {"bid.view", "quot.view"}:
        return redirect("unauthorized")
    rows, counts, f, qs, wards = _filtered(request, pg_rows() if which == "pg" else apg_rows())
    title = "कार्यसम्पादन जमानत (Performance Bank Guarantee)" if which == "pg" else "अग्रिम भुक्तानी जमानत (Advance Payment Guarantee)"
    if request.GET.get("export") == "xlsx":
        return _xlsx(rows, which, title)
    here = request.get_full_path()
    for r in rows:
        r["actions"] = _actions(r, P, here)
    return render(request, "projects/guarantees.html", {
        "which": which, "title": title, "rows": rows, "counts": counts, "f": f, "qs": qs, "statuses": STATUS.items(),
        "total": sum((r["amount"] for r in rows), ZERO), "warn_days": WARN_DAYS,
        "fiscal_years": FiscalYear.objects.all().order_by("-id"),
        "all_wards": sorted_wards(Ward.objects.filter(id__in=wards) if wards is not None else None),
        "show_kind": which == "apg",
        "filter_desc": " · ".join(x for x in (STATUS.get(f["status"], "बहाल (फिर्ता भएका बाहेक)" if not f["status"] else "सबै"),
                                              {"bid": "बोलपत्र", "quot": "दरभाउपत्र"}.get(f["kind"], "")) if x),
    })


def _actions(r, P, here=""):
    from urllib.parse import quote

    from django.urls import reverse

    out = []
    for kind, oid in r["links"]:
        if kind == "bond":
            if "bid.manage" in P and r["status"] != "closed":
                out.append(("म्याद थप", reverse("extend_performance_bond", args=[oid]) + "?next=" + quote(here, safe="/"),
                            False))
            if "bid.print" in P and r["extended"]:
                out.append(("म्याद थप टिप्पणी", reverse("print_performance_bond_tippani", args=[oid]), True))
        elif kind == "bid_advance" and "bid.print" in P:
            out.append(("टिप्पणी", reverse("print_advance_tippani", args=[oid]), True))
        elif kind == "quot_advance" and "quot.print" in P:
            out.append(("टिप्पणी", reverse("print_quatation_advance_tippani", args=[oid]), True))
    wf = "bid_workflow" if r["kind"] == "bid" else "quot_workflow"
    out.append(("कार्यप्रवाह", reverse(wf) + f"?open={r['p'].id}", False))
    return out


def _xlsx(rows, which, title):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "PG" if which == "pg" else "APG"
    head = ["क्र.सं.", "योजनाको नाम", "वडा", "खरिद विधि", "निर्माण व्यवसायी / फर्म", "बैंक", "रकम", "जमानत नम्बर",
            "जारी मिति", "म्याद सकिने मिति", "बाँकी दिन", "स्थिति"]
    ws.append(head)
    for c in ws[1]:
        c.font, c.fill = Font(bold=True), PatternFill("solid", fgColor="E8EEFB")
    for i, r in enumerate(rows, start=1):
        ws.append([i, r["p"].name, getattr(r["p"].ward, "name", ""), r["kind_label"], r["party"], r["bank"],
                   float(r["amount"]), r["number"], r["issue_bs"], r["expiry_bs"], r["days"], r["status_label"]])
    for i, w in enumerate([6, 36, 6, 14, 28, 22, 14, 18, 12, 14, 9, 20], start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{"performance" if which == "pg" else "advance"}-guarantees.xlsx"'
    wb.save(resp)
    return resp


@login_required
def performance_guarantees(request):
    return _page(request, "pg")


@login_required
def advance_guarantees(request):
    return _page(request, "apg")
