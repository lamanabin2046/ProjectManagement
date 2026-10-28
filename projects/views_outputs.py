"""भौतिक उपलब्धि: योजनाको उपलब्धि भर्ने पेज, सूचक सेटिङ, र एकीकृत प्रतिवेदन (Excel सहित)."""
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Count
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils.http import url_has_allowed_host_and_scheme

from accounts.models import Ward
from accounts.permissions import has_perm, user_perms
from accounts.wards import sorted_wards

from .models import FiscalYear, OutputIndicator, Procurement_Type, Project, ProjectOutput, WorkCategory
from .outputs import apply_from_post, catalog, report


def _user_wards(user):
    from usercommittee.views_workflow import _user_wards as uw

    return uw(user)


# ---------------------------------------------------------------- योजनाको उपलब्धि
@login_required
def project_outputs(request, project_id):
    """एउटा योजनाको भौतिक उपलब्धि (अन्तिम टिप्पणी छापिएर मूल्याङ्कन लक भए पनि यहाँबाट सच्याउन मिल्छ)."""
    if not (has_perm(request.user, "uc.valuation") or has_perm(request.user, "bid.manage")):
        return redirect("unauthorized")
    project = get_object_or_404(Project.objects.select_related("ward", "fiscalyear", "work_category"), id=project_id)
    wards = _user_wards(request.user)
    if wards is not None and project.ward_id not in wards:
        return redirect("unauthorized")

    if request.method == "POST":
        errors = apply_from_post(project, request.POST, request.user)
        if errors:
            for e in errors:
                messages.error(request, e)
        else:
            n = project.outputs.count()
            messages.success(request, f"भौतिक उपलब्धि सुरक्षित भयो ({n} सूचक)।")
            nxt = request.POST.get("next") or ""
            if nxt and url_has_allowed_host_and_scheme(nxt, allowed_hosts={request.get_host()}):
                return redirect(nxt)
            return redirect("project_outputs", project_id=project.id)

    return render(request, "projects/project_outputs.html", {
        "project": project, "out_groups": catalog(project), "out_category_id": project.work_category_id,
        "next": request.GET.get("next", ""),
    })


# ---------------------------------------------------------------- सेटिङ: प्रकार र सूचक
@login_required
def output_settings(request):
    if not has_perm(request.user, "settings.outputs"):
        return redirect("unauthorized")

    if request.method == "POST":
        action = request.POST.get("action")
        if action == "add_category":
            name = (request.POST.get("name") or "").strip()[:150]
            if name:
                last = WorkCategory.objects.order_by("-sort").first()
                WorkCategory.objects.create(name=name, sort=(last.sort + 10) if last else 10)
                messages.success(request, f"“{name}” कामको प्रकार थपियो।")
        elif action == "add_indicator":
            cat = get_object_or_404(WorkCategory, id=request.POST.get("category"))
            name = (request.POST.get("name") or "").strip()[:200]
            unit = (request.POST.get("unit") or "").strip()[:40]
            if not name or not unit:
                messages.error(request, "सूचकको नाम र एकाइ दुवै चाहिन्छ।")
            else:
                last = cat.indicators.order_by("-sort").first()
                OutputIndicator.objects.create(category=cat, name=name, unit=unit, sort=(last.sort + 10) if last else 10)
                messages.success(request, f"“{cat.name}” मा “{name} ({unit})” थपियो।")
        elif action == "save":
            for c in WorkCategory.objects.all():
                name = (request.POST.get(f"cat_name_{c.id}") or "").strip()
                if name:
                    c.name = name[:150]
                c.sort = _int(request.POST.get(f"cat_sort_{c.id}"), c.sort)
                c.is_active = bool(request.POST.get(f"cat_active_{c.id}"))
                c.save()
            for i in OutputIndicator.objects.all():
                name = (request.POST.get(f"ind_name_{i.id}") or "").strip()
                unit = (request.POST.get(f"ind_unit_{i.id}") or "").strip()
                if name:
                    i.name = name[:200]
                if unit:
                    i.unit = unit[:40]
                i.sort = _int(request.POST.get(f"ind_sort_{i.id}"), i.sort)
                i.is_active = bool(request.POST.get(f"ind_active_{i.id}"))
                i.save()
            messages.success(request, "सुरक्षित भयो।")
        elif action == "delete_indicator":
            i = get_object_or_404(OutputIndicator, id=request.POST.get("id"))
            if i.values.exists():
                messages.error(request, f"“{i.name}” योजनामा प्रयोग भइसकेकोले हटाउन मिल्दैन — निष्क्रिय गर्नुहोस्।")
            else:
                i.delete()
                messages.success(request, f"“{i.name}” हटाइयो।")
        return redirect("output_settings")

    cats = []
    used = {r["indicator_id"]: r["n"] for r in
            ProjectOutput.objects.values("indicator_id").annotate(n=Count("id"))}
    for c in WorkCategory.objects.prefetch_related("indicators"):
        cats.append({"c": c, "inds": [{"i": i, "used": used.get(i.id, 0)} for i in c.indicators.all()]})
    return render(request, "projects/output_settings.html", {"cats": cats})


def _int(v, default):
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return default


# ---------------------------------------------------------------- एकीकृत प्रतिवेदन
@login_required
def outputs_report(request):
    if not has_perm(request.user, "report.outputs"):
        return redirect("unauthorized")
    qs = Project.objects.all()
    my_wards = _user_wards(request.user)
    if my_wards is not None:
        qs = qs.filter(ward_id__in=my_wards)

    fy = request.GET.get("fiscalyear", "")
    ward = request.GET.get("ward", "")
    ptype = request.GET.get("procurement_type", "")
    cat = request.GET.get("category", "")
    done = request.GET.get("done", "")
    if fy.isdigit():
        qs = qs.filter(fiscalyear_id=fy)
    if ward.isdigit():
        qs = qs.filter(ward_id=ward)
    if ptype.isdigit():
        qs = qs.filter(procurement_type_id=ptype)
    if cat.isdigit():
        qs = qs.filter(outputs__indicator__category_id=cat).distinct()
    if done == "1":   # अन्तिम मूल्याङ्कन भइसकेका (उपभोक्ता समिति) मात्र
        qs = qs.filter(final_payment__isnull=False)

    wards = sorted_wards(Ward.objects.filter(id__in=my_wards) if my_wards is not None else None)
    if ward.isdigit():
        wards = [w for w in wards if str(w.id) == ward]
    groups = report(qs, wards)
    if cat.isdigit():
        groups = [g for g in groups if str(g["cat"].id) == cat]

    if request.GET.get("export") == "xlsx":
        return _xlsx(groups, wards, request)

    fys = FiscalYear.objects.all().order_by("-name")
    return render(request, "projects/outputs_report.html", {
        "groups": groups, "wards": wards,
        "all_wards": sorted_wards(Ward.objects.filter(id__in=my_wards) if my_wards is not None else None),
        "fiscal_years": fys, "procurement_types": Procurement_Type.objects.all(),
        "categories": WorkCategory.objects.all(),
        "f": {"fiscalyear": fy, "ward": ward, "procurement_type": ptype, "category": cat, "done": done},
        "project_count": qs.filter(outputs__isnull=False).distinct().count(),
        "fy_name": next((x.name for x in fys if str(x.id) == fy), ""),
        "qs": request.GET.urlencode(),
    })


def _xlsx(groups, wards, request):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "भौतिक उपलब्धि"
    head = ["कामको प्रकार", "उपलब्धि", "एकाइ", "जम्मा", "योजना संख्या"] + [f"वडा {w.name}" for w in wards]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEFB")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for g in groups:
        for r in g["rows"]:
            ws.append([g["cat"].name, r["ind"].name, r["ind"].unit, float(r["total"]), r["count"]]
                      + [float(v) if v else None for v in r["ward_values"]])
    ws2 = wb.create_sheet("योजना अनुसार")
    ws2.append(["कामको प्रकार", "उपलब्धि", "एकाइ", "योजना", "वडा", "आ.व.", "परिमाण"])
    for c in ws2[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEFB")
    for g in groups:
        for r in g["rows"]:
            for p, q in r["projects"]:
                ws2.append([g["cat"].name, r["ind"].name, r["ind"].unit, p.name,
                            getattr(p.ward, "name", ""), getattr(p.fiscalyear, "name", ""), float(q)])
    for sheet, widths in ((ws, [18, 34, 10, 12, 12] + [10] * len(wards)), (ws2, [18, 34, 10, 40, 8, 10, 12])):
        for i, wdt in enumerate(widths, start=1):
            sheet.column_dimensions[sheet.cell(row=1, column=i).column_letter].width = wdt
        sheet.freeze_panes = "A2"
    resp = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = 'attachment; filename="bhautik-uplabdhi.xlsx"'
    wb.save(resp)
    return resp
